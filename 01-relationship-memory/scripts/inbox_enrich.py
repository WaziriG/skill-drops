"""Relationship Memory: learn each person's email history from an inbox. Read-only on mail, always.  (version 1.1)

Reads your people table (CSV, Excel or JSON export of the spine) and one inbox source:
  --mbox FILE            an exported mailbox (Google Takeout, Apple Mail, Thunderbird)
  --imap-host HOST       a live IMAP connection (folders opened with EXAMINE, messages fetched with BODY.PEEK)
  --messages-json FILE   messages already pulled by Claude through a mail connector (see GUIDE.md)

Works out, per person: Emails Exchanged, Email Threads, First/Last Email, Last Email From, Follow-up Cadence (only when
blank), Mail Search Link and Inbox Checked, plus Warmth and Follow-up Status for reports and spreadsheets. With
--summarize, Claude (`claude -p`) also writes Relationship Context, adds to Open Loops, and proposes a Suggested Can Help
With. It also lists NEW people: anyone with 2+ emails each way who is not in the spine.

Incremental by default: for anyone with stored history and an Inbox Checked date, only mail dated after Inbox Checked
(or after the stored Last Email, if that is later) is read, and it is ADDED to the stored totals. The earlier First Email is kept, Relationship Context is kept and a dated
line is added for what is new, and Open Loops only ever grows. --full recomputes everything from the mail it can see.

It never writes the owner's columns: Name, Email, Other Emails, Company, Role / Title, Segment, Priority, Know Them?,
Can Help With, Notes, Next Step. Follow-up Cadence is written only when it is blank. Warmth and Follow-up Status are
formulas in Notion and Airtable, so they go to a separate "computed" list in --out and are never meant to be written there.

Default is a dry run: it prints what it would change and writes nothing.
  python3 inbox_enrich.py --people people-spine.csv --mbox mock-inbox.mbox --me alex@yourfirm.example
  ... --config ~/.relationship-memory/config.json   addresses, owner, search link, cadence options, column names
  ... --out changes.json        machine-readable changes, for Claude to apply to Notion or Airtable
  ... --apply                   write the changes into the CSV or Excel people file (a backup is kept)
  ... --promote                 move ticked rows from the Discoveries list into the people file
"""
import argparse, csv, datetime as dt, email, getpass, imaplib, json, mailbox, os, re, shutil, statistics, subprocess, sys, time
from concurrent.futures import ThreadPoolExecutor
from email.header import decode_header, make_header
from email.utils import getaddresses, parsedate_to_datetime
from html import unescape
from pathlib import Path
from urllib.parse import quote

VERSION = "1.1"
TODAY = dt.date.today()
OWNER_COLS = ("Name", "Email", "Other Emails", "Company", "Role / Title", "Can Help With", "Know Them?", "Notes",
              "Segment", "Priority", "Next Step")
FORMULA_COLS = ("Warmth", "Follow-up Status")
WRITABLE_COLS = ("Emails Exchanged", "Email Threads", "First Email", "Last Email", "Last Email From", "Relationship Context",
                 "Open Loops", "Follow-up Cadence", "Inbox Checked", "Mail Search Link", "Suggested Can Help With")
MACHINE_COLS = WRITABLE_COLS + FORMULA_COLS
ID_COLS = ("id", "ID", "Page ID", "Record ID", "url", "URL")
DISCOVERY_COLS = ["Promote", "Name", "Email", "Sent", "Received", "First Email", "Last Email", "Last Subject", "Also In", "Found On"]
AUTOMATED = re.compile(r"(no-?reply|do-?not-?reply|notification|newsletter|digest|mailer-daemon|bounce|postmaster|calendar|"
                       r"^news@|^updates?@|^marketing@|^support@|^billing@|^invoices?@|^receipts?@|^hello@|^info@)")
REPLY = re.compile(r"^\s*(re|fw|fwd|aw)\s*:", re.I)
NONE_WORDS = {"", "none", "n/a", "na", "nothing", "no open loops"}
TOUCH_GAP = 21  # emails closer together than this are one conversation, not separate touches
SEARCH_LINKS = {"gmail": "https://mail.google.com/mail/u/0/#search/{query}", "none": ""}
DEFAULT_CADENCE_OPTIONS = {"Every 2 weeks": 14, "Monthly": 30, "Every 6 weeks": 42, "Quarterly": 90, "Twice a year": 182}


# ---------- reading mail ----------

def decode(value):
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return str(value)


def body_text(msg):
    plain, html = [], []
    for part in msg.walk():
        if part.get_content_maintype() != "text" or str(part.get("Content-Disposition", "")).startswith("attachment"):
            continue
        try:
            text = part.get_payload(decode=True).decode(part.get_content_charset() or "utf-8", errors="replace")
        except Exception:
            continue
        (plain if part.get_content_subtype() == "plain" else html).append(text)
    text = "\n".join(plain) or re.sub(r"<[^>]+>", " ", unescape("\n".join(html)))
    kept = []
    for line in text.splitlines():
        if line.startswith(">") or re.match(r"^\s*On .{5,120} wrote:\s*$", line) or line.strip() in ("--", "-- "):
            break
        kept.append(line)
    return re.sub(r"\s+", " ", " ".join(kept)).strip()[:700]


def parse_people(field_values):
    flat = []
    for v in field_values:
        if isinstance(v, (list, tuple)):
            flat += [str(x) for x in v if x]
        elif v:
            flat.append(str(v))
    return [(decode(n).strip().lower(), a.strip().lower()) for n, a in getaddresses([decode(v) for v in flat])]


def to_record(msg, with_body=True, seq=0):
    try:
        stamp = parsedate_to_datetime(msg.get("Date"))
        when = stamp.date()
    except Exception:
        return None
    sender = parse_people(msg.get_all("From", []))
    return {"id": msg.get("Message-ID", "") or f'{msg.get("Date")}{msg.get("Subject")}', "date": when,
            "ts": stamp.isoformat(), "seq": seq, "from": sender[0] if sender else ("", ""),
            "to": parse_people(msg.get_all("To", [])), "cc": parse_people(msg.get_all("Cc", [])),
            "subject": decode(msg.get("Subject", "")).strip(), "text": body_text(msg) if with_body else ""}


def load_mbox(path):
    box = mailbox.mbox(path)
    try:
        return [r for r in (to_record(m, True, i) for i, m in enumerate(box)) if r]
    finally:
        box.close()


def load_messages_json(path):
    """Messages pulled through a connector: [{id, date (YYYY-MM-DD or ISO datetime), from, to, cc, subject, text}]."""
    out = []
    for i, m in enumerate(json.loads(Path(path).read_text())):
        try:
            when = dt.date.fromisoformat(str(m["date"])[:10])
        except Exception:
            continue
        sender = parse_people([m.get("from", "")])
        out.append({"id": m.get("id") or f'{m["date"]}{m.get("subject")}', "date": when, "ts": str(m["date"]), "seq": i,
                    "from": sender[0] if sender else ("", ""), "to": parse_people([m.get("to", "")]),
                    "cc": parse_people([m.get("cc", "")]), "subject": str(m.get("subject", "")).strip(),
                    "text": re.sub(r"\s+", " ", str(m.get("text", "")))[:700]})
    return out


class Imap:
    """Read-only IMAP. Folders are opened read-only (EXAMINE); messages are fetched with BODY.PEEK so nothing is marked read."""

    def __init__(self, host, user, folders, oauth_env=None):
        self.conn = imaplib.IMAP4_SSL(host)
        self.seq = 0
        if oauth_env:
            token = os.environ[oauth_env]
            self.conn.authenticate("XOAUTH2", lambda _: f"user={user}\x01auth=Bearer {token}\x01\x01".encode())
        else:
            password = os.environ.get("INBOX_PASSWORD") or getpass.getpass(f"App password for {user}: ")
            self.conn.login(user, password)
        self.folders = folders or self.default_folders(host)

    def default_folders(self, host):
        listed = " ".join(f.decode(errors="replace") for f in self.conn.list()[1] or [])
        if "gmail" in host or "[Gmail]/All Mail" in listed:
            return ["[Gmail]/All Mail"]
        wanted = ["INBOX", "Sent", "Sent Items", "Sent Messages", "Archive"]
        return [f for f in wanted if f'"{f}"' in listed or f" {f}" in listed] or ["INBOX"]

    def _fetch(self, uids, section, with_body):
        out = []
        for i in range(0, len(uids), 200):
            status, parts = self.conn.uid("FETCH", b",".join(uids[i:i + 200]), f"({section})")
            if status != "OK":
                continue
            for part in parts:
                if isinstance(part, tuple):
                    self.seq += 1
                    rec = to_record(email.message_from_bytes(part[1]), with_body, self.seq)
                    if rec:
                        out.append(rec)
        return out

    def headers_since(self, since):
        """Headers only, every message since a date: cheap enough to scan a week or a quarter."""
        out = []
        for folder in self.folders:
            if self.conn.select(f'"{folder}"', readonly=True)[0] != "OK":
                continue
            status, data = self.conn.uid("SEARCH", None, f'SINCE {since.strftime("%d-%b-%Y")}')
            uids = (data[0] or b"").split() if status == "OK" else []
            out += self._fetch(uids, "BODY.PEEK[HEADER.FIELDS (FROM TO CC DATE SUBJECT MESSAGE-ID)]", False)
        return out

    def history(self, person, cap):
        """History with one person: their most recent `cap` messages per folder, bodies included."""
        terms = [f'OR OR FROM "{a}" TO "{a}" CC "{a}"' for a in person["emails"]] or [f'OR FROM "{person["name"]}" TO "{person["name"]}"']
        out = []
        for folder in self.folders:
            if self.conn.select(f'"{folder}"', readonly=True)[0] != "OK":
                continue
            for term in terms:
                status, data = self.conn.uid("SEARCH", None, f"({term})")
                uids = (data[0] or b"").split()[-cap:] if status == "OK" else []
                out += self._fetch(uids, "BODY.PEEK[]", True)
        return out


# ---------- the people table ----------

def norm_name(name):
    name = re.sub(r",.*$", "", str(name or "")).lower()
    name = re.sub(r"\(.*?\)|[^a-z\s\-']", " ", name)
    return re.sub(r"\s+", " ", name).strip()


def as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return dt.date.fromisoformat(str(value).strip()[:10])
    except Exception:
        return None


def as_int(value):
    """A stored number, or None when the cell is blank (blank is not the same as zero)."""
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def blank(value):
    return value is None or str(value).strip() == ""


def read_table(path, sheet):
    path = Path(path)
    if path.suffix.lower() == ".json":
        rows = json.loads(path.read_text())
        return rows, list(rows[0].keys()) if rows else []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)[sheet]
        head = [c.value for c in ws[1]]
        rows = [{h: ("" if v is None else v) for h, v in zip(head, r) if h} for r in ws.iter_rows(min_row=2, values_only=True)]
        return [r for r in rows if any(str(v).strip() for v in r.values())], [h for h in head if h]
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return list(reader), list(reader.fieldnames or [])


def to_person(row, index):
    name = str(row.get("Name") or f'{row.get("First Name", "")} {row.get("Last Name", "")}').strip()
    raw = f'{row.get("Email") or ""};{row.get("Other Emails") or ""}'
    emails = list(dict.fromkeys(e.strip().lower() for e in re.split(r"[;,\s]+", raw) if "@" in e))
    return {"index": index, "row": row, "name": name, "name_key": norm_name(name), "emails": emails,
            "id": next((str(row[c]) for c in ID_COLS if row.get(c)), ""), "company": str(row.get("Company") or "")}


def stored_view(row, column_map):
    """The row as the script sees it: store column names (e.g. "Gmail History") read back under the schema names."""
    view = dict(row)
    for schema_name, store_name in column_map.items():
        if store_name in row:
            view[schema_name] = row[store_name]
    return view


# ---------- metrics ----------

def index_mail(records):
    by_addr, by_name = {}, {}
    for rec in records:
        for name, addr in [rec["from"]] + rec["to"] + rec.get("cc", []):
            if addr:
                by_addr.setdefault(addr, []).append(rec)
            if name:
                by_name.setdefault(norm_name(name), []).append(rec)
    return by_addr, by_name


def match(person, by_addr, by_name):
    found = {r["id"]: r for a in person["emails"] for r in by_addr.get(a, [])}
    if found:
        return list(found.values()), "Email"
    found = {r["id"]: r for r in by_name.get(person["name_key"], [])} if person["name_key"] else {}
    return list(found.values()), ("Name" if found else "")


def strip_prefix(subject):
    return re.sub(r"^((re|fw|fwd|aw)\s*:\s*)+", "", subject, flags=re.I).strip()


def thread_key(subject):
    return strip_prefix(subject).lower()


def order(recs):
    return sorted({r["id"]: r for r in recs}.values(), key=lambda r: (r["date"], str(r.get("ts", "")), r.get("seq", 0)))


def direction(rec, person, me):
    """Who had the last word in this message, between the owner and this person only.
    Them: sent from one of the person's addresses, with the owner in To.  Me: sent by the owner, with the person in To.
    Anything else (they were only cc'd, a third party wrote, the owner was only cc'd) is None: it counts toward volume,
    never toward who wrote last."""
    sender = rec["from"][1]
    to_addrs = {a for _, a in rec["to"]}
    if person["emails"]:
        is_them = sender in person["emails"]
        them_in_to = bool(set(person["emails"]) & to_addrs)
    else:
        key = person["name_key"]
        is_them = bool(key) and norm_name(rec["from"][0]) == key
        them_in_to = bool(key) and key in {norm_name(n) for n, _ in rec["to"] if n}
    if is_them and me & to_addrs:
        return "Them"
    if sender in me and them_in_to:
        return "Me"
    return None


def metrics(recs, me, person):
    recs = order(recs)
    days = sorted({r["date"] for r in recs})
    touches = [d for i, d in enumerate(days) if i == 0 or (d - days[i - 1]).days > TOUCH_GAP]
    gaps = [(b - a).days for a, b in zip(touches, touches[1:])]
    direct = [(r, d) for r in recs for d in [direction(r, person, me)] if d]
    firsts = {}
    for r in recs:
        firsts.setdefault(thread_key(r["subject"]), r)
    return {"messages": recs, "count": len(recs), "first": recs[0]["date"], "last_any": recs[-1]["date"],
            "last_word": direct[-1][0]["date"] if direct else None, "last_from": direct[-1][1] if direct else None,
            "threads": len(firsts), "thread_firsts": firsts,
            "rhythm": statistics.median(gaps) if len(gaps) >= 2 else None}


def cadence_days(value, options):
    v = str(value or "").strip()
    if not v:
        return None
    try:
        return int(float(v))
    except ValueError:
        pass
    return next((d for label, d in options.items() if label.lower() == v.lower()), None)


def cadence(row, stats, options, select):
    """Returns (days used for status, proposal or None). A cadence the owner already set is never replaced."""
    current = row.get("Follow-up Cadence")
    if not blank(current):
        return cadence_days(current, options) or 90, None
    base = {"A": 30, "B": 60, "C": 90}.get(str(row.get("Priority") or "").strip()[:1].upper(), 90)
    if stats and stats["rhythm"] and stats["rhythm"] < base:
        base = max(14, int(round(stats["rhythm"])))
    if select:
        label = min(options, key=lambda k: (abs(options[k] - base), options[k]))
        return options[label], label
    return base, base


def warmth(count, last, last_from):
    """Volume up to 40 + recency up to 40 + they wrote last 20. 35 or more = warm enough to ask."""
    if not count:
        return 0
    recency = max(0.0, 40 - (TODAY - last).days / 9) if last else 0
    return round(min(40, count * 4) + recency + (20 if last_from == "Them" else 0))


def follow_up_status(last, last_from, cad, know_them=""):
    if not last:
        return "No email history"
    days = (TODAY - last).days
    if last_from == "Them" and days <= 60 and know_them != "No":
        return "Reply owed"
    if days > 3 * cad:
        return "Dormant"
    due = last + dt.timedelta(days=cad)
    if due < TODAY:
        return "Overdue"
    return "Due soon" if (due - TODAY).days <= 14 else "On track"


def search_link(template, person):
    if not template or not person["emails"]:
        return ""
    query = " OR ".join(f"from:{a} OR to:{a}" for a in person["emails"])
    return template.format(query=quote(query), email=quote(person["emails"][0]))


def norm_text(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def append_loops(existing, new, day):
    """Open Loops only grows: a new loop is added on its own dated line unless it is already there. Nothing is removed."""
    new = str(new or "").strip()
    existing = str(existing or "")
    if norm_text(new) in NONE_WORDS or norm_text(new) in norm_text(existing):
        return None
    line = f"{day}: {new}"
    return f"{existing.rstrip()}\n{line}" if existing.strip() else line


def merge_context(existing, new, day):
    """Relationship Context keeps every word already there; what is new is added as a dated line."""
    new = str(new or "").strip()
    if norm_text(new) in NONE_WORDS or norm_text(new) in norm_text(existing):
        return None
    return f"{str(existing).rstrip()}\n{day}: {new}"


# ---------- summaries (claude -p) ----------

PROMPT_FULL = """You are reading the email history between {owner} and {name}{company}. Messages run oldest to newest; "YOU" is {owner}.

{messages}

Return only a JSON object with these keys:
"relationship_context": 2 to 3 plain sentences: how the relationship started, what you have talked about, and anything you worked on together.
"open_loops": anything promised, asked or left unanswered on either side that still needs action, in one sentence, or "None".
"suggested_can_help_with": what {name} could help other people with, judged only from what they say about their own work or what they visibly did in these emails, 25 words or fewer, or "" if the emails do not show it.
Refer to {owner} as "you". Use only what the emails say. Do not guess. Do not mention any other person by name."""

PROMPT_NEW = """You are reading only the NEWEST emails between {owner} and {name}{company}, those dated after {cutoff}. "YOU" is {owner}.

What is already recorded about this relationship (do not repeat it):
{context}

Open loops already recorded (do not repeat them):
{loops}

New messages, oldest to newest:
{messages}

Return only a JSON object with these keys:
"whats_new": one or two plain sentences on what these new emails add to the relationship, or "" if nothing of substance.
"open_loops": anything NEW that was promised, asked or left unanswered in these emails and still needs action, in one sentence, or "None".
"suggested_can_help_with": what {name} could help other people with, judged only from what they say about their own work or what they visibly did in these emails, 25 words or fewer, or "".
Refer to {owner} as "you". Use only what the emails say. Do not guess. Do not mention any other person by name."""


def summarize(person, msgs, owner, model, me, existing=None, cutoff=None):
    lines = []
    for r in msgs[-15:]:
        if r["from"][1] in me:
            who = "YOU"
        elif direction(r, person, me) == "Them" or r["from"][1] in person["emails"]:
            who = person["name"]
        else:
            who = f'{r["from"][0] or r["from"][1]} (someone else; {person["name"]} copied)'
        lines.append(f'[{r["date"]}] {who}: {r["subject"]} :: {r["text"][:500]}')
    company = f' ({person["company"]})' if person["company"] else ""
    if existing is not None:
        prompt = PROMPT_NEW.format(owner=owner, name=person["name"], company=company, cutoff=cutoff, messages="\n".join(lines),
                                   context=existing.get("context") or "(nothing yet)", loops=existing.get("loops") or "(none)")
    else:
        prompt = PROMPT_FULL.format(owner=owner, name=person["name"], company=company, messages="\n".join(lines))
    error = ""
    for attempt in range(3):
        try:
            run = subprocess.run(["claude", "-p", "--model", model, "--tools", "", "--no-session-persistence", "--strict-mcp-config"],
                                 input=prompt, capture_output=True, text=True, timeout=180)
            out = run.stdout
            data = json.loads(out[out.index("{"): out.rindex("}") + 1])
            return {k: str(v).strip() for k, v in data.items()}
        except Exception as exc:
            error = (run.stderr.strip() if "run" in locals() and run.stderr else str(exc))[:200]
            time.sleep(4 * (attempt + 1))
    return {"_error": error}


# ---------- the core: who changed and what to write ----------

def plan(person, recs, how, me, opts):
    """One person's proposed change before summaries, or None when there is nothing new to say."""
    row = person["row"]
    checked = as_date(row.get("Inbox Checked"))
    stored_n = as_int(row.get("Emails Exchanged"))
    has_history = stored_n is not None or not blank(row.get("Last Email")) or not blank(row.get("Last Email From"))
    incremental = not opts["full"] and checked is not None and has_history
    # Stored totals already cover mail up to Inbox Checked, or up to Last Email if the store recorded a later one.
    cutoff = max(d for d in (checked, as_date(row.get("Last Email"))) if d) if incremental else None
    window = [r for r in recs if r["date"] > cutoff] if incremental else list(recs)
    know = str(row.get("Know Them?") or "").strip()
    base = {"index": person["index"], "id": person["id"], "name": person["name"], "email": (person["emails"] or [""])[0],
            "match": how, "mode": "incremental" if incremental else "full", "new_messages": len(window)}

    if not window:
        if incremental:
            return None  # no new mail since Inbox Checked: stored history stands as is
        if stored_n:  # full run found nothing, but the store has history: never zero it on an empty result
            return {**base, "kind": "kept", "fields": {}, "computed": {}}
        if str(row.get("Last Email From") or "") == "No email yet" and checked:
            return None
        return {**base, "kind": "no-mail",
                "fields": {"Inbox Checked": TODAY.isoformat(), "Emails Exchanged": 0, "Last Email From": "No email yet"},
                "computed": {"Warmth": 0, "Follow-up Status": "No email history"}}

    s = metrics(window, me, person)
    f = {"Inbox Checked": TODAY.isoformat()}
    link = search_link(opts["template"], person)
    if link:
        f["Mail Search Link"] = link
    if incremental:
        older_keys = {thread_key(r["subject"]) for r in recs if r["date"] <= cutoff}
        new_threads = [k for k, r in s["thread_firsts"].items() if k not in older_keys and not REPLY.match(r["subject"])]
        f["Emails Exchanged"] = (stored_n or 0) + s["count"]
        f["Email Threads"] = (as_int(row.get("Email Threads")) or 0) + len(new_threads)
        if not as_date(row.get("First Email")):
            f["First Email"] = min(r["date"] for r in recs).isoformat()
        stored_last = as_date(row.get("Last Email"))
        if s["last_word"] and (not stored_last or s["last_word"] >= stored_last):
            f["Last Email"], f["Last Email From"] = s["last_word"].isoformat(), s["last_from"]
        elif not stored_last:
            f["Last Email"] = s["last_any"].isoformat()
    else:
        f.update({"Emails Exchanged": s["count"], "Email Threads": s["threads"], "First Email": s["first"].isoformat(),
                  "Last Email": (s["last_word"] or s["last_any"]).isoformat(), "Last Email From": s["last_from"] or ""})

    cad, proposal = cadence(row, s, opts["cadence_options"], opts["cadence_select"])
    if proposal is not None:
        f["Follow-up Cadence"] = proposal
    if blank(row.get("Relationship Context")):
        subjects = list(dict.fromkeys(strip_prefix(m["subject"]) for m in reversed(s["messages"])))
        f["Relationship Context"] = "Email subjects so far: " + "; ".join(subjects[:5])

    n = f.get("Emails Exchanged", stored_n or 0)
    last = as_date(f.get("Last Email")) or as_date(row.get("Last Email"))
    last_from = f.get("Last Email From", row.get("Last Email From") or "")
    computed = {"Warmth": warmth(int(n), last, last_from) if last else 0,
                "Follow-up Status": follow_up_status(last, last_from, cad, know)}
    return {**base, "kind": "history", "fields": f, "computed": computed, "stats": s, "cutoff": cutoff}


def fold_summary(change, row, summary):
    """Apply Claude's summary under the 1.1 rules: context merges, loops append, suggestions only where the owner is silent."""
    f, day = change["fields"], TODAY.isoformat()
    existing_ctx = str(row.get("Relationship Context") or "").strip()
    if change["mode"] == "incremental" and existing_ctx:
        merged = merge_context(existing_ctx, summary.get("whats_new", ""), day)
        if merged:
            f["Relationship Context"] = merged
    else:
        ctx = summary.get("relationship_context") or summary.get("whats_new") or ""
        if ctx:
            f["Relationship Context"] = ctx
    loops = append_loops(row.get("Open Loops"), summary.get("open_loops", ""), day)
    if loops:
        f["Open Loops"] = loops
    chw = summary.get("suggested_can_help_with", "")
    if chw and blank(row.get("Can Help With")) and str(row.get("Know Them?") or "").strip() != "No":
        f["Suggested Can Help With"] = chw


def build_changes(people, found, me, opts, summarizer=None):
    """people: to_person() list. found: {index: (records, how)}; how None = skipped (IMAP, no activity).
    summarizer(person, msgs, existing, cutoff) -> dict; None turns summaries off."""
    changes, todo = [], []
    for p in people:
        recs, how = found.get(p["index"], ([], None))
        if how is None and p["index"] in found:
            continue
        c = plan(p, recs, how, me, opts)
        if c is None:
            continue
        changes.append(c)
        if c["kind"] == "history":
            todo.append(c)
    failed = []
    if summarizer and todo:
        by_index = {p["index"]: p for p in people}

        def one(c):
            p = by_index[c["index"]]
            row = p["row"]
            existing = None
            if c["mode"] == "incremental" and not blank(row.get("Relationship Context")):
                existing = {"context": str(row.get("Relationship Context")), "loops": str(row.get("Open Loops") or "")}
            return summarizer(p, c["stats"]["messages"], existing, c["cutoff"])

        with ThreadPoolExecutor(max_workers=4) as pool:
            for c, s in zip(todo, pool.map(one, todo)):
                if "_error" in s:
                    failed.append((c["name"], s["_error"]))
                else:
                    fold_summary(c, by_index[c["index"]]["row"], s)
    for c in changes:
        c.pop("stats", None)
        bad = set(c["fields"]) - set(WRITABLE_COLS)
        assert not bad, f"refusing to write non-history columns for {c['name']}: {bad}"
        assert not set(c["computed"]) - set(FORMULA_COLS)
    return changes, failed


def diff_for(change, row):
    both = {**change["fields"], **change["computed"]}
    return {k: (row.get(k, ""), v) for k, v in both.items() if str(row.get(k, "") or "") != str(v)}


# ---------- new people ----------

def discover(records, me, known_addrs, known_names, feeder_addrs, feeder_names, min_each_way):
    sent, received, info = {}, {}, {}
    for r in records:
        if r["from"][1] in me:
            pairs = [(n, a) for n, a in r["to"] + r.get("cc", []) if a and a not in me]
            for n, a in pairs:
                sent[a] = sent.get(a, 0) + 1
        else:
            pairs = [r["from"]] if r["from"][1] else []
            for n, a in pairs:
                received[a] = received.get(a, 0) + 1
        for n, a in pairs:
            i = info.setdefault(a, {"name": "", "first": r["date"], "last": r["date"], "subject": r["subject"]})
            i["name"] = i["name"] or n
            i["first"], i["last"] = min(i["first"], r["date"]), max(i["last"], r["date"])
            if r["date"] >= i["last"]:
                i["subject"] = r["subject"]
    out = []
    for a, i in info.items():
        if sent.get(a, 0) < min_each_way or received.get(a, 0) < min_each_way or AUTOMATED.search(a):
            continue
        if a in known_addrs or (i["name"] and norm_name(i["name"]) in known_names):
            continue
        name = i["name"].title() if i["name"] else a
        also = "LinkedIn feeder" if a in feeder_addrs or norm_name(name) in feeder_names else ""
        out.append({"Promote": "", "Name": name, "Email": a, "Sent": sent[a], "Received": received[a],
                    "First Email": i["first"].isoformat(), "Last Email": i["last"].isoformat(),
                    "Last Subject": strip_prefix(i["subject"]), "Also In": also, "Found On": TODAY.isoformat()})
    return sorted(out, key=lambda d: (-(d["Sent"] + d["Received"]), d["Name"]))


# ---------- writing (CSV / Excel only; Notion and Airtable are written by Claude from --out) ----------

def backup(path):
    archive = path.parent / "_archive"
    archive.mkdir(exist_ok=True)
    shutil.copy2(path, archive / f"{path.stem}-before-{dt.datetime.now():%Y%m%d-%H%M%S}{path.suffix}")


def apply_csv(path, rows, head, changes, discoveries):
    head = head + [c for c in MACHINE_COLS if c not in head]
    for ch in changes:
        rows[ch["index"]].update({k: v for k, v in {**ch["fields"], **ch["computed"]}.items() if k in MACHINE_COLS})
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, head, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    merge_discoveries_csv(path.with_name(path.stem + "-discoveries.csv"), discoveries)


def merge_discoveries_csv(disc_path, discoveries):
    existing = read_table(disc_path, None)[0] if disc_path.exists() else []
    seen = {r["Email"] for r in existing}
    rows = existing + [d for d in discoveries if d["Email"] not in seen]
    with open(disc_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, DISCOVERY_COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def apply_xlsx(path, sheet, changes, discoveries):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    wb = load_workbook(path)
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    table = next(iter(ws.tables.values()), None)
    for c in MACHINE_COLS:
        if c not in head:
            head.append(c)
            cell = ws.cell(row=1, column=len(head), value=c)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A3F6B")
            if table is not None:
                table.tableColumns.append(TableColumn(id=len(table.tableColumns) + 1, name=c))
    if table is not None:
        first = table.ref.split(":")[0]
        table.ref = f"{first}:{get_column_letter(len(head))}{ws.max_row}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    col = {h: get_column_letter(i + 1) for i, h in enumerate(head) if h}
    for ch in changes:
        r = ch["index"] + 2
        for k, v in {**ch["fields"], **ch["computed"]}.items():
            if k not in MACHINE_COLS:
                continue
            cell = ws[f"{col[k]}{r}"]
            if k == "Warmth":
                n, le, lf = col["Emails Exchanged"], col["Last Email"], col["Last Email From"]
                v = (f'=IF(N({n}{r})=0,0,ROUND(MIN(40,{n}{r}*4)+IF({le}{r}="",0,MAX(0,40-(TODAY()-{le}{r})/9))'
                     f'+IF({lf}{r}="Them",20,0),0))')
            elif k == "Follow-up Status":
                le, lf, cd = col["Last Email"], col["Last Email From"], col["Follow-up Cadence"]
                not_no = f',{col["Know Them?"]}{r}<>"No"' if "Know Them?" in col else ""
                v = (f'=IF({le}{r}="","No email history",IF(AND({lf}{r}="Them",TODAY()-{le}{r}<=60{not_no}),"Reply owed",'
                     f'IF(TODAY()-{le}{r}>3*{cd}{r},"Dormant",IF({le}{r}+{cd}{r}<TODAY(),"Overdue",'
                     f'IF({le}{r}+{cd}{r}-TODAY()<=14,"Due soon","On track")))))')
            elif k in ("First Email", "Last Email", "Inbox Checked") and v:
                v = dt.date.fromisoformat(v)
                cell.number_format = "yyyy-mm-dd"
            cell.value = v
    ds = wb["Discoveries"] if "Discoveries" in wb.sheetnames else wb.create_sheet("Discoveries")
    if ds.max_row < 1 or ds.cell(row=1, column=1).value != "Promote":
        ds.delete_rows(1, ds.max_row)
        ds.append(DISCOVERY_COLS)
    seen = {ds.cell(row=i, column=3).value for i in range(2, ds.max_row + 1)}
    for d in discoveries:
        if d["Email"] not in seen:
            ds.append([d[c] for c in DISCOVERY_COLS])
    wb.save(path)


def promote(path, sheet):
    """Move ticked Discoveries rows into the people file with Know Them? left blank for the owner to decide."""
    ticked = lambda v: str(v or "").strip().lower() in ("y", "yes", "x", "true", "1")
    note = f"Added from inbox discovery {TODAY.isoformat()}"
    backup(path)
    if path.suffix.lower() == ".csv":
        disc_path = path.with_name(path.stem + "-discoveries.csv")
        if not disc_path.exists():
            return 0
        disc = read_table(disc_path, None)[0]
        rows, head = read_table(path, None)
        moving = [d for d in disc if ticked(d.get("Promote"))]
        rows += [{**{h: "" for h in head}, "Name": d["Name"], "Email": d["Email"], "Notes": note} for d in moving]
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, head, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        with open(disc_path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, DISCOVERY_COLS, extrasaction="ignore")
            w.writeheader()
            w.writerows([d for d in disc if not ticked(d.get("Promote"))])
        return len(moving)
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path)
    if "Discoveries" not in wb.sheetnames:
        return 0
    ws, ds = wb[sheet], wb["Discoveries"]
    head = [c.value for c in ws[1]]
    moved = 0
    for i in range(ds.max_row, 1, -1):
        if ticked(ds.cell(row=i, column=1).value):
            new = {"Name": ds.cell(row=i, column=2).value, "Email": ds.cell(row=i, column=3).value, "Notes": note}
            ws.append([new.get(h, None) for h in head])
            ds.delete_rows(i)
            moved += 1
    table = next(iter(ws.tables.values()), None)
    if table is not None:
        table.ref = f'{table.ref.split(":")[0]}:{get_column_letter(len(head))}{ws.max_row}'
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    wb.save(path)
    return moved


# ---------- report ----------

def short(v, n=110):
    s = "" if v is None else str(v)
    return s if len(s) <= n else s[: n - 1] + "…"


def report(args, people, changes, discoveries, name_only, failed, source_desc, n_records, rows):
    mode = "APPLY" if args.apply else "DRY RUN (nothing written)"
    lines = [f"RELATIONSHIP MEMORY {VERSION}: inbox learning run, {TODAY.isoformat()} - {mode}",
             f"Inbox: {source_desc} ({n_records} messages read)   People in spine: {len(people)}", ""]
    updated = [c for c in changes if c["kind"] == "history"]
    owed = [c for c in updated if c["computed"].get("Follow-up Status") == "Reply owed"]
    incr = sum(1 for c in updated if c["mode"] == "incremental")
    kept = [c["name"] for c in changes if c["kind"] == "kept"]
    lines.append(f"History updated: {len(updated)} ({incr} added to stored totals, {len(updated) - incr} full)   "
                 f"No email yet: {sum(1 for c in changes if c['kind'] == 'no-mail')}   "
                 f"Unchanged since last check: {len(people) - len(changes)}   Replies owed: {len(owed)}")
    if owed:
        lines.append("Replies owed: " + ", ".join(c["name"] for c in owed))
    if kept:
        lines.append("Stored history kept (this run found no mail for them; check the source): " + ", ".join(kept))
    lines += ["", "PROPOSED CHANGES (history columns only; owner columns, including Next Step, are never touched)"]
    for c in changes:
        d = diff_for(c, rows[c["index"]])
        if not d:
            continue
        tag = f' [matched by {c["match"]}]' if c["match"] else ""
        tag += f' [+{c["new_messages"]} new since last check]' if c["mode"] == "incremental" else ""
        lines.append(f'- {c["name"]}{tag}')
        for k, (old, new) in d.items():
            if k in ("Inbox Checked", "Mail Search Link"):
                continue
            lines.append(f"    {k}: {short(old) or '(blank)'} -> {short(new)}")
    sugg = [c for c in changes if c["fields"].get("Suggested Can Help With")]
    if sugg:
        lines += ["", "SUGGESTED CAN HELP WITH (for the owner to accept or ignore; never copied into Can Help With)"]
        for c in sugg:
            lines.append(f'- {c["name"]}: "{c["fields"]["Suggested Can Help With"]}"')
    lines += ["", f"NEW PEOPLE ({args.min_each_way}+ emails each way, not in the spine): {len(discoveries)}"]
    for d in discoveries:
        lines.append(f'- {d["Name"]} <{d["Email"]}>: {d["Sent"]} sent / {d["Received"]} received, last {d["Last Email"]}, '
                     f'"{d["Last Subject"]}"' + (f' (also in {d["Also In"]})' if d["Also In"] else ""))
    if name_only:
        lines += ["", "CHECK THESE (matched by name only, no email address on file): " + ", ".join(name_only)]
    if failed:
        lines += ["", f"{len(failed)} summaries failed; run again with --summarize to fill only those. First error: {failed[0][1]}"]
    return "\n".join(lines)


# ---------- main ----------

def load_config(path):
    if not path:
        return {}
    p = Path(path).expanduser()
    return json.loads(p.read_text()) if p.exists() else {}


def store_setting(config, key, legacy_key=None):
    """A setting at the top of config.json, or inside the store's own section (e.g. config["notion"])."""
    if key in config:
        return config[key]
    section = config.get(config.get("store", ""), {}) or {}
    return section.get(key, section.get(legacy_key) if legacy_key else None)


def parse_cadence_options(text):
    out = {}
    for part in str(text or "").split(","):
        if "=" in part:
            label, days = part.rsplit("=", 1)
            out[label.strip()] = int(days)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--people", required=True, help="Spine people table: .csv, .xlsx (sheet --sheet) or .json export")
    ap.add_argument("--sheet", default="People", help="Excel sheet holding the spine (default People)")
    ap.add_argument("--feeder", action="append", default=[], help="Optional feeder table(s), used to label new people")
    ap.add_argument("--config", help="config.json from setup: my_addresses, owner_name, search_link, cadence options, column_map")
    src = ap.add_mutually_exclusive_group()
    src.add_argument("--mbox", help="Exported mailbox (.mbox)")
    src.add_argument("--imap-host", help="imap.gmail.com, outlook.office365.com, imap.mail.me.com, ...")
    src.add_argument("--messages-json", help="Messages pulled by Claude through a mail connector")
    ap.add_argument("--imap-user", help="IMAP login")
    ap.add_argument("--imap-folders", help="Comma-separated folders (default: All Mail on Gmail, else Inbox + Sent)")
    ap.add_argument("--imap-oauth-env", help="Name of an env var holding an OAuth2 access token (XOAUTH2, e.g. Microsoft 365)")
    ap.add_argument("--cap", type=int, default=200, help="Most recent messages per person per folder (IMAP)")
    ap.add_argument("--scan-days", type=int, default=120, help="IMAP: days of headers to scan for activity and new people")
    ap.add_argument("--since", help="Ignore mail before this date (YYYY-MM-DD)")
    ap.add_argument("--me", default="", help="Your own address(es), comma-separated")
    ap.add_argument("--owner", default=None, help="Your name, used in summaries")
    ap.add_argument("--full", action="store_true", help="Recompute everyone from all the mail it can see, instead of adding new mail to stored totals")
    ap.add_argument("--summarize", action="store_true", help="Have Claude write Relationship Context, Open Loops and Suggested Can Help With")
    ap.add_argument("--model", default=None, help="Claude model for summaries (default sonnet)")
    ap.add_argument("--search-link", default=None, help="gmail, none, or a template containing {query} or {email}")
    ap.add_argument("--cadence-options", help='Follow-up Cadence is a select: "Every 2 weeks=14,Monthly=30,..." (or "default")')
    ap.add_argument("--min-each-way", type=int, default=2, help="New people need at least this many emails each way (default 2)")
    ap.add_argument("--out", help="Write the changes as JSON (Claude applies these to Notion or Airtable)")
    ap.add_argument("--report", help="Also save the printed report to this file")
    ap.add_argument("--apply", action="store_true", help="Write into the CSV or Excel people file (backup kept in _archive/)")
    ap.add_argument("--promote", action="store_true", help="Move ticked Discoveries rows into the people file, then stop")
    args = ap.parse_args()

    people_path = Path(args.people).expanduser()
    if args.promote:
        print(f"Promoted {promote(people_path, args.sheet)} people into {people_path}. Know Them? is left blank for you.")
        return
    if not (args.mbox or args.imap_host or args.messages_json):
        sys.exit("Give one inbox source: --mbox, --imap-host or --messages-json.")

    config = load_config(args.config)
    owner = args.owner or config.get("owner_name") or "the inbox owner"
    model = args.model or config.get("summary_model") or "sonnet"
    link_kind = args.search_link or config.get("search_link") or "none"
    column_map = store_setting(config, "column_map") or {}
    cad_cfg = store_setting(config, "follow_up_cadence_options", "follow_up_cadence_is_select")
    if args.cadence_options:
        cad_cfg = DEFAULT_CADENCE_OPTIONS if args.cadence_options == "default" else parse_cadence_options(args.cadence_options)
    opts = {"full": args.full, "template": SEARCH_LINKS.get(link_kind, link_kind),
            "cadence_options": cad_cfg or DEFAULT_CADENCE_OPTIONS, "cadence_select": bool(cad_cfg)}

    raw_rows, head = read_table(people_path, args.sheet)
    rows = [stored_view(r, column_map) for r in raw_rows]
    people = [to_person(r, i) for i, r in enumerate(rows)]
    me = {a.strip().lower() for a in args.me.split(",") if a.strip()} or {a.lower() for a in config.get("my_addresses", [])}
    if args.imap_user:
        me.add(args.imap_user.lower())
    if not me:
        sys.exit("Give your own address with --me (or my_addresses in --config) so sent and received mail can be told apart.")
    since = as_date(args.since) if args.since else None

    # 1. read mail
    found, records = {}, []
    if args.imap_host:
        if not args.imap_user:
            sys.exit("--imap-user is required with --imap-host")
        imap = Imap(args.imap_host, args.imap_user, [f.strip() for f in (args.imap_folders or "").split(",") if f.strip()], args.imap_oauth_env)
        scan_from = max(since or dt.date.min, TODAY - dt.timedelta(days=args.scan_days))
        records = imap.headers_since(scan_from)
        active_addrs, active_names = index_mail(records)
        for p in people:
            checked = as_date(p["row"].get("Inbox Checked"))
            recent = match(p, active_addrs, active_names)[0]
            if args.full or not checked or checked < scan_from or any(r["date"] > checked for r in recent):
                recs = imap.history(p, args.cap)
                found[p["index"]] = (recs, ("Email" if p["emails"] else "Name") if recs else "")
            else:
                found[p["index"]] = ([], None)
        source_desc = f"IMAP {args.imap_host}"
    else:
        records = load_mbox(args.mbox) if args.mbox else load_messages_json(args.messages_json)
        by_addr, by_name = index_mail(records)
        for p in people:
            found[p["index"]] = match(p, by_addr, by_name)
        source_desc = f"mbox {Path(args.mbox).name}" if args.mbox else f"messages {Path(args.messages_json).name}"
    if since:
        records = [r for r in records if r["date"] >= since]
        found = {k: ([r for r in v[0] if r["date"] >= since], v[1]) for k, v in found.items()}

    # 2-3. decide who changed, compute history, summaries
    summarizer = None
    if args.summarize:
        summarizer = lambda p, msgs, existing, cutoff: summarize(p, msgs, owner, model, me, existing, cutoff)
    changes, failed = build_changes(people, found, me, opts, summarizer)

    # 4. new people
    feeders = [to_person(r, i) for f in args.feeder for i, r in enumerate(read_table(Path(f).expanduser(), args.sheet)[0])]
    discoveries = discover(records, me, {a for p in people for a in p["emails"]}, {p["name_key"] for p in people},
                           {a for p in feeders for a in p["emails"]}, {p["name_key"] for p in feeders}, args.min_each_way)
    name_only = [c["name"] for c in changes if c["match"] == "Name"]

    text = report(args, people, changes, discoveries, name_only, failed, source_desc, len(records), rows)
    print(text)
    if args.report:
        Path(args.report).write_text(text + "\n")
    if args.out:
        rename = lambda fields: {column_map.get(k, k): v for k, v in fields.items()}
        Path(args.out).write_text(json.dumps({
            "run": {"date": TODAY.isoformat(), "version": VERSION, "source": source_desc, "applied": bool(args.apply),
                    "people": len(people), "full": bool(args.full)},
            "owner_columns_never_written": list(OWNER_COLS),
            "formulas_never_written": list(FORMULA_COLS),
            "rules": {"Open Loops": "append-only: the value given already contains every earlier line",
                      "Relationship Context": "incremental runs keep the stored text and add a dated line",
                      "Follow-up Cadence": "present only when the stored value was blank"},
            "updates": [{**{k: c[k] for k in ("id", "name", "email", "match", "kind", "mode", "new_messages")},
                         "fields": rename(c["fields"]), "computed": c["computed"]}
                        for c in changes if c["kind"] != "kept"],
            "kept_without_change": [c["name"] for c in changes if c["kind"] == "kept"],
            "new_people": discoveries, "name_only_matches": name_only,
            "summary_failures": [{"name": n, "error": e} for n, e in failed]}, indent=1, default=str))
    if args.apply:
        if people_path.suffix.lower() not in (".csv", ".xlsx", ".xlsm"):
            sys.exit("--apply writes CSV or Excel files only. For Notion or Airtable, use --out and let Claude apply it.")
        live = [c for c in changes if c["kind"] != "kept"]
        backup(people_path)
        if people_path.suffix.lower() == ".csv":
            apply_csv(people_path, raw_rows, head, live, discoveries)
        else:
            apply_xlsx(people_path, args.sheet, live, discoveries)
        print(f"\nWrote {len(live)} rows to {people_path} (backup in {people_path.parent / '_archive'}).")


if __name__ == "__main__":
    main()
