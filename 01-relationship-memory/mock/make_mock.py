"""Build the demo data for the Relationship Memory Kit. Every person, company, address and email is invented.

Writes, next to this file:
  people-spine.csv       30 people in the spine table, owner columns filled, history columns empty
  People.xlsx            the same spine as an Excel table (sheet "People"), for the Excel route
  linkedin-feeder.csv    40 LinkedIn-style feeder rows (some overlap the spine)
  mock-inbox.mbox        an inbox with real back-and-forth, replies owed, open loops, new people and noise

Mail dates are set relative to the day you run this, so the demo always looks current.
  python3 make_mock.py
"""
import csv, datetime as dt, mailbox, random
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from pathlib import Path

HERE = Path(__file__).resolve().parent
ME = ("Alex Rivera", "alex@yourfirm.example")
NOW = dt.datetime.combine(dt.date.today(), dt.time(9, 0), tzinfo=dt.timezone(dt.timedelta(hours=-5)))
random.seed(1001)

OWNER_COLS = ["Name", "Email", "Company", "Role / Title", "Segment", "Priority", "Know Them?", "Can Help With", "Notes"]
MACHINE_COLS = ["Suggested Can Help With", "Emails Exchanged", "Email Threads", "First Email", "Last Email",
                "Last Email From", "Relationship Context", "Open Loops", "Next Step", "Follow-up Cadence",
                "Inbox Checked", "Mail Search Link", "Warmth", "Follow-up Status"]

# name, email, company, role, segment, priority, know, can help with (owner's words), notes
SPINE = [
    ("Avery Calderwood", "avery@ashgrovecapital.example", "Ashgrove Capital Partners", "Managing Partner", "Investor", "A", "Yes",
     "Invests in B2B software at seed and Series A, checks of $250K to $1M. Knows most of the regional angel groups.", "Met at a founder dinner."),
    ("Benedikt Fenwick", "ben@latticerow.example", "Lattice Row Accounting", "Founder & CPA", "Partner", "B", "Yes",
     "Fractional CFO work for service firms under $20M revenue. Clean books, cash forecasting, getting ready for a bank loan.", ""),
    ("Corinne Hazelton", "corinne.hazelton@tallgrasshealth.example", "Tallgrass Health", "Chief Operating Officer", "Client", "A", "Yes",
     "Runs operations for a 12-clinic healthcare group. Knows HIPAA compliance cold and can talk vendor selection for patient scheduling.", ""),
    ("Dayo Northcote", "dayo@riverstonetalent.example", "Riverstone Talent", "Principal Recruiter", "Partner", "B", "Yes",
     "Places senior operations and finance hires. Fast on COO and controller searches in the Northeast.", ""),
    ("Emeka Stanhope", "emeka@brightmanlegal.example", "Brightman Legal", "Partner", "Peer", "B", "Yes",
     "Commercial contracts and SaaS agreements. Good first call before signing a vendor or partnership agreement.", ""),
    ("Freya Redfern", "freya@kestrelcyber.example", "Kestrel Cyber", "Founder", "Peer", "A", "Yes",
     "Incident response and tabletop exercises for mid-size firms. Also knows the cyber insurance underwriters.", ""),
    ("Gustavo Mallory", "gmallory@copperline.example", "Copperline Logistics", "VP, Operations", "Client", "B", "Yes",
     "Warehouse and fleet operations. Has rolled out route-planning software across 40 trucks.", ""),
    ("Hiroshi Vandermeer", "hiroshi@northfieldvc.example", "Northfield Ventures", "Principal", "Investor", "B", "Yes",
     "", "Can Help With left blank on purpose: the weekly run should suggest one."),
    ("Ines Kittredge", "ines@inglebyevents.example", "Ingleby Events", "Owner", "Partner", "C", "Yes",
     "Puts the right 30 people in a room. Runs a monthly operators dinner and makes introductions generously.", ""),
    ("Jarrah Elsworth", "jelsworth@quarrylanebank.example", "Quarry Lane Bank", "SVP, Commercial Lending", "Peer", "B", "Yes",
     "Commercial lending and SBA loans for companies between $2M and $50M revenue.", ""),
    ("Kalani Oakhurst", "kalani@wexfordsoftware.example", "Wexford Software", "Chief Technology Officer", "Client", "A", "Yes",
     "Hiring and managing offshore engineering teams. Built a data platform on Snowflake.", ""),
    ("Leopold Thistlewood", "leopold@thistlewoodyarrow.example", "Thistlewood & Yarrow", "Managing Director", "Prospect", "B", "Yes",
     "M&A advisory for owner-led businesses selling for $5M to $40M.", ""),
    ("Marisa Underhill", "marisa@pennanthr.example", "Pennant HR Partners", "Founder", "Partner", "B", "Yes",
     "Outsourced HR for teams of 10 to 150: handbooks, onboarding, and tricky terminations.", ""),
    ("Nkem Dunmore", "nkem@galbraithfoods.example", "Galbraith Foods", "Chief Executive Officer", "Client", "A", "Yes",
     "Food manufacturing and co-packing. Knows how to get a product onto regional grocery shelves.", ""),
    ("Orla Jessop", "orla@larkspurmedia.example", "Larkspur Media", "Creative Director", "Vendor", "C", "Yes",
     "Brand identity and short-form video. Fast turnaround on launch videos.", ""),
    ("Pascal Brightwater", "pascal@summitlineins.example", "Summitline Insurance", "Commercial Broker", "Peer", "C", "Yes",
     "Commercial insurance, including cyber and key-person policies.", ""),
    ("Quinn Penhallow", "quinn@penhallowadvisory.example", "Penhallow Advisory", "Independent Consultant", "Peer", "C", "Yes",
     "Pricing strategy for professional services firms. Has moved three firms from hourly to fixed fee.", ""),
    ("Rania Zabriskie", "rania@cobaltgrove.example", "Cobalt Grove Foundation", "Executive Director", "Partner", "B", "Yes",
     "Grant funding for workforce programs. Sits on two nonprofit boards and knows the city's foundation officers.", ""),
    ("Soren Ledbetter", "soren@ledbettersystems.example", "Ledbetter Systems", "IT Director", "Peer", "C", "Yes",
     "Microsoft 365 administration and device management for 200 seats.", ""),
    ("Thandiwe Marchbank", "thandiwe@marchbankrealty.example", "Marchbank Realty", "Commercial Broker", "Peer", "C", "Yes",
     "", "Blank on purpose: the weekly run should suggest one."),
    ("Ulrich Fairhaven", "ulrich@harbornine.example", "Harbor Nine Capital", "Partner", "Investor", "B", "Getting to know", "", "Met once at a conference panel."),
    ("Valentina Crane", "valentina@cranesupply.example", "Crane Supply Co.", "Owner", "Prospect", "B", "Getting to know", "", ""),
    ("Wren Holbrook", "wren@silverbirchlabs.example", "Silverbirch Labs", "Head of AI", "Prospect", "C", "Getting to know", "", ""),
    ("Ximena Rourke", "ximena@rourkedental.example", "Rourke Dental Group", "Practice Owner", "Prospect", "C", "Getting to know", "", ""),
    ("Yusuf Tarrant", "yusuf@tarrantfreight.example", "Tarrant Freight", "Founder", "Prospect", "C", "Getting to know", "", ""),
    ("Zainab Merrow", "zainab@merrowinteriors.example", "Merrow Interiors", "Principal", "Prospect", "C", "Getting to know", "", ""),
    ("Anders Pike", "anders@pikeleadgen.example", "Pike Lead Gen", "Sales Representative", "Not a Fit", "", "No", "", "Cold pitches only."),
    ("Bryony Clegg", "bryony@cleggprint.example", "Clegg Print", "Account Manager", "Vendor", "", "No", "", "Former vendor. Do not suggest."),
    ("Cosmo Delancey", "cosmo@delanceystudio.example", "Delancey Studio", "Photographer", "Vendor", "", "", "", ""),
    ("Delia Fontaine", "", "Fontaine Wealth", "Wealth Advisor", "Peer", "", "", "", "No email on file: the inbox run should find her by name."),
]

# Threads: (person email or alias, subject, days ago the thread started, days between messages, [(who, text)])
DELIA = "delia.fontaine@fontainewealth.example"
T = [
    ("avery@ashgrovecapital.example", "Following up from the founder dinner", 420, 3, [
        ("me", "Avery, good to meet you at the dinner. You mentioned you look at B2B software at seed. Happy to send over two companies I think you'd like."),
        ("them", "Please do. We're writing seed checks again this quarter."),
        ("me", "Sent both decks. The second one is the one I'd look at first."),
        ("them", "Took the call with the second one. Strong team. We're doing diligence now.")]),
    ("avery@ashgrovecapital.example", "Quarterly catch-up", 45, 4, [
        ("them", "Want to catch up before the holidays? Curious what you're seeing in operations-heavy companies."),
        ("me", "Yes. Thursday the 14th works. I'll bring notes on the three firms we mapped this fall."),
        ("them", "Great conversation. I'll send the intro to our portfolio COO group this week."),
        ("me", "Thank you. I'll prepare a short session for them.")]),
    ("ben@latticerow.example", "Referral: a client who needs cash forecasting", 200, 5, [
        ("me", "Ben, one of our clients needs a real 13-week cash forecast before a bank meeting. Would you take a look?"),
        ("them", "Happy to. Send their contact over and I'll set up a call."),
        ("them", "Finished the forecast. They walked into the bank meeting with it and got the line of credit."),
        ("me", "That's a great outcome. Thank you for moving so fast.")]),
    ("ben@latticerow.example", "Lunch?", 70, 6, [
        ("them", "Lunch next month? I have two clients who need help automating their month-end close."),
        ("me", "Yes, the first week works. Send me what their close looks like today.")]),
    ("corinne.hazelton@tallgrasshealth.example", "Scheduling workflow map", 300, 7, [
        ("me", "Corinne, attached is the map of the patient scheduling workflow across all 12 clinics."),
        ("them", "This is the first time I've seen it on one page. The no-show follow-up is where we lose the most."),
        ("me", "Agreed. We'll start there. Build plan to follow."),
        ("them", "Approved. Let's start on the 3rd.")]),
    ("corinne.hazelton@tallgrasshealth.example", "HR coverage while we hire", 30, 3, [
        ("them", "Our HR manager is leaving at the end of the month and the search will take a while. Do you know a firm that can cover outsourced HR for six months?"),
        ("me", "I do. Let me check with them first and come back to you this week."),
        ("them", "Thank you. Onboarding for the two new clinics is the part I'm most worried about.")]),
    ("dayo@riverstonetalent.example", "Controller search", 260, 4, [
        ("them", "Thanks for the referral last spring. That COO placement worked out well."),
        ("me", "Glad to hear it. I'll keep sending people your way."),
        ("them", "Please do. We have capacity for two more finance searches this quarter.")]),
    ("emeka@brightmanlegal.example", "Partnership agreement review", 150, 5, [
        ("me", "Emeka, could you look at a partnership agreement before we sign? Mostly the IP clause."),
        ("them", "Sent back redlines. The IP clause needed work; the rest is standard."),
        ("me", "Very helpful. Signed with your changes.")]),
    ("freya@kestrelcyber.example", "Tabletop exercise for a client", 120, 4, [
        ("me", "Freya, a client of ours wants a tabletop exercise before their insurance renewal. Interested?"),
        ("them", "Yes. We run those in half a day. Send me their size and stack."),
        ("me", "About 80 people, Microsoft 365, one office. Intro sent."),
        ("them", "Ran the tabletop last week. They found two gaps in their backup restore. Thanks for the referral.")]),
    ("freya@kestrelcyber.example", "Co-hosting a breakfast", 20, 2, [
        ("them", "Want to co-host a breakfast for operators on incident response? I'll bring the content if you bring the room."),
        ("me", "I like it. Let me find a date in the next six weeks.")]),
    ("gmallory@copperline.example", "Route planning rollout", 340, 8, [
        ("me", "Gustavo, here's the rollout plan for route planning across the fleet."),
        ("them", "Looks right. Drivers will push back on the tablets, so training matters."),
        ("me", "Agreed. We'll do it depot by depot."),
        ("them", "All 40 trucks are live. Fuel spend is down 9 percent in the first month.")]),
    ("gmallory@copperline.example", "Thinking ahead", 95, 4, [
        ("them", "Off the record: our owner is starting to think about selling the business in two years. Do you know anyone who advises on that?"),
        ("me", "I do. Let me think about who fits best and come back to you.")]),
    ("hiroshi@northfieldvc.example", "Logistics tech intro", 180, 5, [
        ("me", "Hiroshi, I know a logistics software founder raising a Series A. Worth an intro?"),
        ("them", "Yes. I lead our climate and logistics tech investments, mostly Series A, and I sit on two portfolio boards. Send it."),
        ("me", "Intro sent."),
        ("them", "Met them. Not quite ready, but I'll stay close.")]),
    ("hiroshi@northfieldvc.example", "Portfolio operations help", 12, 3, [
        ("me", "Hiroshi, I saw two of your portfolio companies are scaling operations. Happy to share what we're seeing."),
        ("them", "Yes please. Both are struggling with their reporting. Could you send a short note I can forward to the CEOs?")]),
    ("ines@inglebyevents.example", "Operators dinner", 230, 6, [
        ("them", "You're on the list for next month's operators dinner. Thirty people, no pitches."),
        ("me", "Thank you. I'll be there."),
        ("them", "Thanks for coming. Three people asked me for your details afterwards.")]),
    ("ines@inglebyevents.example", "Next dinner: speaker?", 8, 2, [
        ("them", "Would you speak for ten minutes at the next dinner? Topic is up to you."),
        ("me", "Yes, happy to. Can I share a short story about how operators spend their week?"),
        ("them", "Perfect. Can you send a title and a two-line description by Friday?")]),
    ("jelsworth@quarrylanebank.example", "Intro from a mutual client", 380, 5, [
        ("them", "A mutual client said you helped them get their numbers in order before our meeting. It made the loan easy."),
        ("me", "Glad it helped. Happy to send more companies your way when the timing is right."),
        ("them", "Please do. We like operators who know their numbers.")]),
    ("kalani@wexfordsoftware.example", "Data platform phase two", 160, 6, [
        ("me", "Kalani, phase two scope attached: the reporting layer on top of the new data platform."),
        ("them", "Approved. Let's start on the 1st."),
        ("me", "Phase two delivered. Dashboards are live for the leadership team."),
        ("them", "Leadership loves them. Board meeting went smoothly.")]),
    ("kalani@wexfordsoftware.example", "Two things for Q1", 10, 2, [
        ("me", "Kalani, checking in before planning season. What's on your list for Q1?"),
        ("them", "Two things. Our controller is leaving, so we need to replace them by Q1. And the board wants a cyber tabletop exercise before our insurance renewal. Know anyone for either?")]),
    ("leopold@thistlewoodyarrow.example", "Sell-side workflow", 60, 5, [
        ("them", "We spend too many hours on data rooms for each deal. Curious whether you've mapped that kind of work."),
        ("me", "We have. Happy to walk you through it. Next Tuesday?"),
        ("them", "Tuesday works.")]),
    ("marisa@pennanthr.example", "Handbook project", 210, 5, [
        ("them", "Thanks for the referral. We finished the handbook and onboarding for your client."),
        ("me", "They told me it was the smoothest part of the year."),
        ("them", "Always happy to take more teams in the 10 to 150 range.")]),
    ("nkem@galbraithfoods.example", "Co-packing workflow", 500, 9, [
        ("me", "Nkem, here's the map of the co-packing order workflow."),
        ("them", "This explains why our lead times slip. Let's fix the handoff to the plant."),
        ("me", "Done. The order handoff is automated as of today."),
        ("them", "Lead times are down from 11 days to 6.")]),
    ("nkem@galbraithfoods.example", "Equipment loan and the books", 35, 4, [
        ("them", "We want to refinance the equipment loan this year and I don't have a banker I trust. Also our books need a real CFO's eyes before we talk to anyone."),
        ("me", "Understood. I know people for both. Let me come back to you."),
        ("them", "Thank you. No rush, but before year end would be ideal.")]),
    ("orla@larkspurmedia.example", "Launch video", 310, 6, [
        ("me", "Orla, we need a 60-second launch video by the 20th."),
        ("them", "Doable. Script by Friday, first cut by the 15th."),
        ("them", "Final cut attached.")]),
    ("pascal@summitlineins.example", "Key-person policy", 640, 6, [
        ("me", "Pascal, a client is asking about key-person insurance. Can you help?"),
        ("them", "Yes. Send them over and I'll quote it.")]),
    ("quinn@penhallowadvisory.example", "Fixed-fee pricing", 110, 4, [
        ("them", "Read your note on fixed-fee engagements. We should compare notes."),
        ("me", "Agreed. Coffee next week?"),
        ("them", "Next week works. I'll bring the pricing model I use.")]),
    ("rania@cobaltgrove.example", "Workforce program evaluation", 140, 7, [
        ("me", "Rania, here's the evaluation framework for the workforce program."),
        ("them", "The board loved it. Our funders want this kind of reporting every quarter."),
        ("me", "We'll set it up to run automatically.")]),
    ("soren@ledbettersystems.example", "Device management", 75, 3, [
        ("them", "Saw your talk. We manage 200 seats on Microsoft 365 and I'd like to compare notes on device policies."),
        ("me", "Happy to. Thursday?"),
        ("them", "Thursday works.")]),
    ("thandiwe@marchbankrealty.example", "Office space", 260, 5, [
        ("me", "Thandiwe, we're looking for a small office downtown, about 1,500 square feet."),
        ("them", "I handle office leases downtown and in the business district. I'll send three options."),
        ("them", "Here are three options. The second has the best lease terms."),
        ("me", "We signed the second one. Thank you for negotiating the free months.")]),
    ("ulrich@harbornine.example", "Good to meet you", 50, 3, [
        ("me", "Ulrich, good to meet you on the panel. I'd like to learn more about what Harbor Nine looks for."),
        ("them", "Likewise. Let's find time next month.")]),
    ("valentina@cranesupply.example", "Your workshop", 40, 2, [
        ("them", "I attended your workshop. We're a 40-person distributor and our ordering is still on paper."),
        ("me", "Thanks for coming. Happy to look at how orders move today. Free next week?")]),
    ("wren@silverbirchlabs.example", "Agents in production", 25, 3, [
        ("me", "Wren, enjoyed your post on agents in production. Would you compare notes sometime?"),
        ("them", "Sure, in a few weeks once our launch is out.")]),
    ("ximena@rourkedental.example", "Scheduling question", 90, 2, [
        ("them", "A colleague mentioned you helped a clinic group with scheduling. Is that something you'd do for a dental practice?")]),
    ("yusuf@tarrantfreight.example", "Intro via a mutual friend", 15, 2, [
        ("me", "Yusuf, a mutual friend suggested we talk about dispatch software."),
        ("them", "Happy to. Let's talk after our busy season.")]),
    ("zainab@merrowinteriors.example", "Follow-up", 140, 2, [
        ("me", "Zainab, following up on our chat at the trade show.")]),
    ("anders@pikeleadgen.example", "10x your pipeline", 30, 4, [
        ("them", "Want 50 qualified meetings next month?"),
        ("them", "Bumping this to the top of your inbox."),
        ("them", "Last try: 50 meetings, guaranteed.")]),
    ("bryony@cleggprint.example", "Invoice", 700, 3, [
        ("them", "Attached is the final invoice."),
        ("me", "Paid, thank you.")]),
    (DELIA, "Retirement plan options for your team", 220, 5, [
        ("them", "As promised, here are three retirement plan options for a team your size."),
        ("me", "Thanks Delia. The second option looks right."),
        ("them", "Great. I'll send the paperwork.")]),
    # New people: real two-way back-and-forth, not in the spine
    ("gideon@fernhollow.example", "Workshop collaboration", 65, 4, [
        ("them", "Loved your workshop. Would you co-teach a session for our members?"),
        ("me", "I'd like that. What size group?"),
        ("them", "About 25 owners. Late next month?"),
        ("me", "Works. I'll draft an outline."),
        ("them", "Outline looks great. Confirmed for the 28th."),
        ("me", "Locked in. See you then.")]),
    ("m.lowell@ostranderpartners.example", "Board advisor question", 55, 5, [
        ("them", "We're adding an operations advisor to our board. Would you be open to a conversation?"),
        ("me", "Open to it. What does the role involve?"),
        ("them", "Four meetings a year and some help on the operating plan."),
        ("me", "That's manageable. Let's talk next week.")]),
    ("tobias@ketteringco.example", "Reporting automation", 28, 3, [
        ("them", "A friend sent me your article on Monday reporting. We do this by hand every week."),
        ("me", "Happy to help. What tools do you use today?"),
        ("them", "Spreadsheets and a lot of copy and paste."),
        ("me", "Let's look at it together. Here are two times."),
        ("them", "Tuesday at 10 works.")]),
    # Not new people: too little back-and-forth, or automated
    ("harriet@sloanecreative.example", "Quick hello", 18, 2, [
        ("them", "Hi, we met briefly at the conference."), ("me", "Good to hear from you, Harriet.")]),
]
NAMES = {p[1]: p[0] for p in SPINE if p[1]}
NAMES.update({DELIA: "Delia Fontaine", "gideon@fernhollow.example": "Gideon Ashcombe",
              "m.lowell@ostranderpartners.example": "Marguerite Lowell", "tobias@ketteringco.example": "Tobias Kettering",
              "harriet@sloanecreative.example": "Harriet Sloane"})

FIRST = "Abel Briar Celia Dorian Elspeth Fintan Gemma Hollis Isolde Jasper Keturah Lorcan Maren Nolan Ophira Percival Rosalind Silas Tamsin Ulla Vaughn Wilhelmina Yara Zephyr".split()
LAST = "Ainsworth Blythe Carrow Dewhurst Ellery Fairlie Gresham Haverford Ilsley Jellicoe Kerridge Loxley Moberly Nettleton Ormsby Prideaux Quennell Rushworth Sedley Tolland".split()
FEEDER_ROLES = [("Harborview Credit Union", "VP, Commercial Lending", "Financial Services"), ("Meta", "Product Manager", "Technology"),
                ("Ridgeback Search Partners", "Search Fund Principal (ETA)", "Financial Services"), ("Bellwether Clinics", "Practice Manager", "Healthcare"),
                ("Fernhill Robotics", "Head of Operations", "Manufacturing"), ("Copse & Vale", "Brand Strategist", "Media"),
                ("Graystone Cloud", "Security Engineer", "Technology"), ("Pinecrest Academy", "Head of School", "Education"),
                ("Lowmoor Freight", "Dispatch Manager", "Logistics"), ("Starling Grocers", "Category Buyer", "Retail"),
                ("Whitlock Capital", "Associate", "Financial Services"), ("Meadowlark HR", "HR Business Partner", "Professional Services")]


def addr(name, mail):
    return f"{name} <{mail}>"


def write_mbox():
    msgs = []
    for mail, subject, start_ago, gap, thread in T:
        when, parent = NOW - dt.timedelta(days=start_ago), None
        name = NAMES[mail]
        for n, (who, text) in enumerate(thread):
            m = EmailMessage()
            sender, rcpt = ((ME[0], ME[1]), (name, mail)) if who == "me" else ((name, mail), (ME[0], ME[1]))
            m["From"], m["To"] = addr(*sender), addr(*rcpt)
            m["Subject"] = ("Re: " if n else "") + subject
            m["Date"] = format_datetime(min(when, NOW - dt.timedelta(hours=2)))
            mid = make_msgid(domain="mock.example")
            m["Message-ID"] = mid
            if parent:
                m["In-Reply-To"] = parent
            m.set_content(text + "\n\n-- \n" + sender[0])
            msgs.append(m)
            parent = mid
            when += dt.timedelta(days=gap, hours=random.randint(1, 9))
    for k in range(8):  # newsletter noise: automated, one-way
        m = EmailMessage()
        m["From"], m["To"] = "Weekly Digest <newsletter@digest.example>", addr(*ME)
        m["Subject"], m["Date"] = f"This week in operations #{k + 1}", format_datetime(NOW - dt.timedelta(days=7 * k + 1))
        m["Message-ID"] = make_msgid(domain="digest.example")
        m.set_content("Five links worth your time.")
        msgs.append(m)
    msgs.sort(key=lambda m: dt.datetime.strptime(m["Date"][5:25], "%d %b %Y %H:%M:%S"))
    out = HERE / "mock-inbox.mbox"
    if out.exists():
        out.unlink()
    box = mailbox.mbox(str(out))
    for m in msgs:
        box.add(m)
    box.flush()
    return len(msgs)


def write_spine():
    rows = [dict(zip(OWNER_COLS, p)) for p in SPINE]
    for r in rows:
        r.update({c: "" for c in MACHINE_COLS})
    cols = OWNER_COLS + MACHINE_COLS
    with open(HERE / "people-spine.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, cols)
        w.writeheader()
        w.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        return len(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B2A4A" if c in OWNER_COLS else "2A3F6B")
        ws.column_dimensions[cell.column_letter].width = 40 if c in ("Can Help With", "Relationship Context", "Open Loops", "Next Step", "Suggested Can Help With") else 16
    table = Table(displayName="People", ref=f"A1:{ws.cell(row=1, column=len(cols)).column_letter}{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B2"
    wb.save(HERE / "People.xlsx")
    return len(rows)


def write_feeder():
    rows, used = [], set()
    overlap = [("Avery Calderwood", "Ashgrove Capital Partners", "Managing Partner", "avery@ashgrovecapital.example"),
               ("Kalani Oakhurst", "Wexford Software", "Chief Technology Officer", ""),
               ("Ines Kittredge", "Ingleby Events", "Owner", ""),
               ("Gideon Ashcombe", "Fernhollow Collective", "Community Director", "")]
    for name, co, role, mail in overlap:
        rows.append({"Name": name, "Company": co, "Role / Title": role, "Email": mail, "In Spine": "Yes" if name != "Gideon Ashcombe" else ""})
        used.add(name)
    while len(rows) < 40:
        name = f"{random.choice(FIRST)} {random.choice(LAST)}"
        if name in used:
            continue
        used.add(name)
        co, role, _ = random.choice(FEEDER_ROLES)
        rows.append({"Name": name, "Company": co, "Role / Title": role, "Email": "", "In Spine": ""})
    for r in rows:
        slug = r["Name"].lower().replace(" ", "-")
        r["LinkedIn URL"] = f"https://www.linkedin.example/in/{slug}-{random.randint(100, 999)}"
        r["Connected On"] = (dt.date.today() - dt.timedelta(days=random.randint(200, 2400))).isoformat()
        r["Last Message Date"] = (dt.date.today() - dt.timedelta(days=random.randint(30, 1500))).isoformat()
        r["Messages"] = random.randint(2, 30)
    cols = ["Name", "Company", "Role / Title", "Email", "LinkedIn URL", "Connected On", "Last Message Date", "Messages", "In Spine"]
    with open(HERE / "linkedin-feeder.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, cols)
        w.writeheader()
        w.writerows(rows)
    return len(rows)


if __name__ == "__main__":
    print(f"{write_spine()} spine people, {write_feeder()} feeder rows, {write_mbox()} messages -> {HERE}")
